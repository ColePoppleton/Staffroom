"""Read and write tables as CSV or XLSX.

No pupil data
-------------
Every table read here passes through
:func:`staffroom.core.privacy.drop_identifier_columns`. Columns that could
identify a pupil are removed before the data reaches a tool, a preview or an
output; the names of the dropped columns are kept on ``Table.dropped_columns``
and logged as a warning. There is no way to opt out.

A tool may need to know which row carried a flag word (``marks`` looks for a
``MAX`` row) even when that word sits in an identifier column, as it does in
a markbook export. ``read_table(path, flag="MAX")`` decides that on the raw
table, before the filter runs, and hands the tool only the row indices on
``Table.flagged_rows``; the filter itself is unchanged and the identifier
cells never leave this module.

Input mistakes
--------------
A wrong kind of file, a missing file, a CSV that is really a binary file or a
damaged workbook are the teacher's everyday mistakes, not bugs.
:func:`read_table` turns each into a :class:`~staffroom.core.plugin.ToolError`
whose one-line message names the file and says what would have worked, so
every tool, the CLI and the launcher show a sentence rather than a traceback.
Anything else that escapes is a genuine bug and stays one.

Fake-data marker
----------------
Real pupil data must never be committed. Every example file carries a marker
that the pre-commit hook looks for:

* CSV: the first line is ``# FAKE DATA`` (anything after the words is ignored).
* XLSX: some cell in the first few rows of any sheet contains ``FAKE DATA``.

:func:`read_table` strips the marker and records it on ``Table.fake_data``;
:func:`write_table` puts it back, so derived outputs stay committable.
"""

from __future__ import annotations

import csv
import zipfile
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from staffroom.core.log import get_logger
from staffroom.core.plugin import ToolError
from staffroom.core.privacy import drop_identifier_columns
from staffroom.core.table import Row, Table, coerce

log = get_logger(__name__)

FAKE_DATA_MARKER = "# FAKE DATA"
_MARKER_WORDS = "FAKE DATA"
_MARKER_SCAN_ROWS = 3

TABLE_EXTENSIONS = (".csv", ".xlsx")


def is_marker_line(line: str) -> bool:
    return line.strip().upper().startswith(FAKE_DATA_MARKER)


def read_table(path: Path, *, flag: str | None = None) -> Table:
    """Read a CSV or XLSX file, dispatching on suffix. Identifier columns are
    dropped (see module docstring). With ``flag``, the indices of rows where any
    cell equals it (case-insensitive, whitespace ignored) are recorded on
    ``Table.flagged_rows``; the decision is made before the columns are dropped.

    Raises :class:`ToolError` with a one-line, teacher-facing message for the
    wrong kind of file, a missing or unreadable file, a CSV that is not text
    and a damaged or mis-named workbook."""
    suffix = path.suffix.lower()
    if suffix not in TABLE_EXTENSIONS:
        raise ToolError(unsupported_table_message(path))
    try:
        raw = read_csv(path) if suffix == ".csv" else read_xlsx(path)
    except FileNotFoundError:
        raise ToolError(f"Could not find {path.name} (looked in {path.parent}).") from None
    except OSError as exc:
        reason = exc.strerror or exc.__class__.__name__
        raise ToolError(f"Could not read {path.name}: {reason}.") from exc
    except UnicodeDecodeError as exc:
        raise ToolError(
            f"Could not read {path.name} as text: it is not a plain CSV file. "
            "Save it from Excel or Sheets as CSV (UTF-8), or pick the .xlsx instead."
        ) from exc
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError) as exc:
        # openpyxl's complaints about a file that is not (or no longer) a workbook.
        # A CSV reader raising one of these is a bug, not a damaged spreadsheet.
        if suffix != ".xlsx":
            raise
        raise ToolError(
            f"Could not open {path.name} as a spreadsheet: it is damaged or not really "
            "an Excel file. Re-save it from Excel as .xlsx, or export a CSV."
        ) from exc
    flagged = flagged_rows(raw, flag) if flag else []
    clean = _privacy_filter(raw, path)
    clean.flagged_rows = flagged
    return clean


def unsupported_table_message(path: Path) -> str:
    """One line for a file whose type no table reader accepts."""
    kind = f"a {path.suffix} file" if path.suffix else "a file with no file extension"
    accepted = " or ".join(TABLE_EXTENSIONS)
    return f"{path.name} is {kind}: choose a {accepted} table."


def flagged_rows(table: Table, flag: str) -> list[int]:
    """Indices of rows where any cell is the word ``flag`` (case-insensitive)."""
    want = flag.strip().upper()
    return [
        i
        for i, row in enumerate(table.rows)
        if any(isinstance(c, str) and c.strip().upper() == want for c in row)
    ]


def _privacy_filter(table: Table, path: Path) -> Table:
    clean, dropped = drop_identifier_columns(table)
    if dropped:
        log.warning("%s: dropped identifier column(s) %s", path.name, ", ".join(dropped))
    return clean


def write_table(table: Table, path: Path) -> Path:
    """Write ``table`` as CSV or XLSX by extension, creating the folder first."""
    suffix = path.suffix.lower()
    path.parent.mkdir(parents=True, exist_ok=True)
    if suffix == ".csv":
        return write_csv(table, path)
    if suffix == ".xlsx":
        return write_xlsx(table, path)
    msg = f"unsupported table file {path.name!r}; expected one of {TABLE_EXTENSIONS}"
    raise ValueError(msg)


# ---------------------------------------------------------------- CSV
def read_csv(path: Path) -> Table:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        lines = fh.read().splitlines()
    fake = False
    while lines and (is_marker_line(lines[0]) or not lines[0].strip()):
        fake = fake or is_marker_line(lines[0])
        lines.pop(0)
    if not lines:
        return Table([], [], fake_data=fake)
    reader = csv.reader(lines)
    header = [h.strip() for h in next(reader)]
    rows: list[Row] = [[coerce(c) for c in _pad(r, len(header))] for r in reader if any(r)]
    return Table(header, rows, fake_data=fake)


def write_csv(table: Table, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        if table.fake_data:
            fh.write(FAKE_DATA_MARKER + "\n")
        writer = csv.writer(fh)
        writer.writerow(table.columns)
        for row in table.rows:
            writer.writerow(["" if c is None else c for c in row])
    return path


# ---------------------------------------------------------------- XLSX
def read_xlsx(path: Path, sheet: str | None = None) -> Table:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    fake = xlsx_has_marker(path)
    # Drop marker/blank rows before the header.
    while raw and (_row_is_marker(raw[0]) or not any(c is not None for c in raw[0])):
        raw.pop(0)
    if not raw:
        return Table([], [], fake_data=fake)
    header = ["" if c is None else str(c).strip() for c in raw[0]]
    rows: list[Row] = [
        [coerce(c) for c in _pad(r, len(header))] for r in raw[1:] if any(c is not None for c in r)
    ]
    return Table(header, rows, fake_data=fake)


def write_xlsx(table: Table, path: Path, sheet: str = "Sheet1") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet
    if table.fake_data:
        ws.append([FAKE_DATA_MARKER])
    ws.append(table.columns)
    for row in table.rows:
        ws.append(row)
    wb.save(path)
    return path


def xlsx_has_marker(path: Path) -> bool:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _MARKER_SCAN_ROWS:
                    break
                if _row_is_marker(list(row)):
                    return True
    finally:
        wb.close()
    return False


def csv_has_marker(path: Path) -> bool:
    with path.open(encoding="utf-8-sig", errors="replace") as fh:
        for _ in range(_MARKER_SCAN_ROWS):
            line = fh.readline()
            if is_marker_line(line):
                return True
    return False


def has_fake_data_marker(path: Path) -> bool:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return csv_has_marker(path)
    if suffix == ".xlsx":
        return xlsx_has_marker(path)
    return False


# ---------------------------------------------------------------- helpers
def _pad(row: Sequence[object], width: int) -> list[object]:
    padded = list(row[:width])
    padded.extend([None] * (width - len(padded)))
    return padded


def _row_is_marker(row: Sequence[object]) -> bool:
    return any(isinstance(c, str) and _MARKER_WORDS in c.upper() for c in row)
