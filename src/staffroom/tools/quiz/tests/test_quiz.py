"""Bank parsing, exporters, and the acceptance for the retrieval scheduler:
it picks only from topics taught in the configured windows."""

import csv
import io
import logging
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from staffroom.cli import app
from staffroom.core.config import load_config
from staffroom.core.csvio import read_table
from staffroom.core.plugin import RunContext, ToolError
from staffroom.core.table import Table
from staffroom.tools.quiz import logic, plugin
from staffroom.tools.quiz.export import write_kahoot
from staffroom.tools.quiz.logic import KAHOOT_HEADERS, Options

runner = CliRunner()
EXAMPLES = Path(__file__).resolve().parents[5] / "examples"


@pytest.fixture
def bank() -> logic.Bank:
    return logic.parse_bank((EXAMPLES / "question_bank.md").read_text())


def test_parse_example_bank(bank: logic.Bank) -> None:
    assert bank.errors == []
    assert len(bank.questions) == 13
    q = bank.questions[0]
    assert q.qtype == "mcq" and q.correct == (0,) and q.options[0] == "Central Processing Unit"
    assert q.tags == ("ocr:1.1.1", "aqa:3.4.4", "topic:cpu") and q.topics == ("cpu",)
    assert q.why.startswith("The CPU")
    tf = bank.questions[2]
    assert tf.qtype == "tf" and tf.options == ("True", "False") and tf.correct == (0,)
    short = bank.questions[4]
    assert short.qtype == "short" and short.answer == "10"
    assert bank.topics == ["cpu", "memory", "binary", "loops", "searching", "networks"]


def test_parse_errors_have_line_numbers() -> None:
    text = (
        "stray line\n## No options\ntype: mcq\n## Bad tf\ntype: tf\nanswer: maybe\n"
        "## Fine\nanswer: 4\n"
    )
    b = logic.parse_bank(text)
    assert [q.text for q in b.questions] == ["Fine"]
    assert any("line 1" in e for e in b.errors)
    assert any("line 2" in e and "two" in e for e in b.errors)
    assert any("line 4" in e and "true" in e for e in b.errors)


def test_filters_and_choose(bank: logic.Bank) -> None:
    binary = logic.filter_questions(bank.questions, topics="binary")
    assert {q.number for q in binary} == {5, 6, 7, 8}
    memory_unit = logic.filter_questions(bank.questions, spec="ocr:1.2")
    assert {q.number for q in memory_unit} == {3, 4, 5, 6, 7, 8}
    assert logic.filter_questions(bank.questions, spec="aqa:3.3.2") == binary[:3]
    chosen = logic.choose(bank.questions, 5, seed=3)
    assert len(chosen) == 5 and chosen == sorted(chosen, key=lambda q: q.number)
    assert logic.choose(bank.questions, 5, seed=3) == chosen


def test_retrieval_picks_only_from_taught_windows(bank: logic.Bank) -> None:
    entries, errors = logic.parse_taught_log(read_table(EXAMPLES / "taught_log.csv"))
    assert errors == [] and len(entries) == 6
    today = date(2026, 10, 8)
    windows = logic.topics_in_windows(entries, today)
    assert windows == {1: ["networks"], 3: ["loops"], 6: ["cpu"]}
    chosen, _ = logic.retrieval_pick(bank.questions, entries, today, 10, seed=1)
    assert chosen
    assert all(set(q.topics) & {"networks", "loops", "cpu"} for q in chosen)
    assert not any("memory" in q.topics or "binary" in q.topics for q in chosen)
    assert {q.topics[0] for q in chosen} == {"networks", "loops", "cpu"}
    # A date with nothing taught in any window picks nothing.
    none, w = logic.retrieval_pick(bank.questions, entries, date(2027, 1, 1), 10, seed=1)
    assert none == [] and all(v == [] for v in w.values())
    bad, errs = logic.parse_taught_log(
        Table(["date", "topic"], [["not-a-date", "cpu"], ["2026-01-01", None]])
    )
    assert bad == [] and errs == [
        "line 2: bad date 'not-a-date' (use YYYY-MM-DD)",
        "line 3: missing topic",
    ]
    assert logic.parse_taught_log(Table(["x", "y"], [[1, 2]]))[1] == [
        "taught log needs 'date' and 'topic' columns"
    ]
    assert logic.parse_taught_log(Table([]))[1] == ["taught log is empty"]
    # A date cell from an .xlsx log arrives with a midnight time.
    excel, errs = logic.parse_taught_log(Table(["Date", "Topic"], [["2026-09-17 00:00:00", "CPU"]]))
    assert errs == [] and excel == [logic.TaughtEntry(date(2026, 9, 17), "cpu")]


def test_taught_log_warns_about_extra_columns_but_never_their_values(tmp_path: Path) -> None:
    # A column the log should not have (#32): it is ignored, named once in the
    # messages, and nothing in it is ever repeated back.
    path = tmp_path / "taught.csv"
    path.write_text(
        "# FAKE DATA\n"
        "Date,Room,topic,Period,\n"
        "2026-09-17,IT2,loops,P3,\n"
        "2026-09-24,IT2,searching,P3,\n",
        encoding="utf-8",
    )
    entries, messages = logic.parse_taught_log(read_table(path))
    assert [e.topic for e in entries] == ["loops", "searching"]
    assert messages == [
        "ignored column 'Room' — the log is date and topic only",
        "ignored column 'Period' — the log is date and topic only",
        "ignored unnamed column 5 — the log is date and topic only",
    ]
    joined = " ".join(messages)
    assert "IT2" not in joined and "P3" not in joined
    # The usual log has nothing to say.
    assert logic.parse_taught_log(Table(["date", "topic"], [["2026-09-17", "loops"]]))[1] == []


def test_plugin_logs_extra_taught_log_columns(tmp_path: Path) -> None:
    log_path = tmp_path / "taught.csv"
    log_path.write_text("# FAKE DATA\ndate,topic,room\n2026-10-01,networks,IT4\n", encoding="utf-8")
    logged: list[str] = []
    ctx = RunContext(config=load_config(), output_dir=tmp_path / "out", log=logged.append)
    options = Options(fmt="forms", retrieval=True, today="2026-10-08", count=3, seed=1)
    result = plugin.run(
        {"source": EXAMPLES / "question_bank.md", "taught_log": log_path, "options": options},
        ctx,
    )
    assert result.tables["chosen"].rows
    assert "Log: ignored column 'room' — the log is date and topic only" in logged
    assert not any("IT4" in m for m in logged)


class _Records(logging.Handler):
    def __init__(self) -> None:
        super().__init__(logging.DEBUG)
        self.messages: list[str] = []

    def emit(self, record: logging.LogRecord) -> None:
        self.messages.append(record.getMessage())


def test_taught_log_identifier_columns_are_dropped_by_csvio(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # #68: the taught log goes through core.csvio like every other table, so a
    # column that could identify a pupil is dropped before the quiz sees it.
    # The cells are placeholders, never names, real or invented.
    log_path = tmp_path / "taught.csv"
    log_path.write_text(
        "# FAKE DATA\n"
        "date,name,topic,Class\n"
        "2026-09-17,CELL-A1,loops,CELL-B1\n"
        "2026-10-01,CELL-A2,networks,CELL-B2\n",
        encoding="utf-8",
    )
    seen: list[Table] = []
    parse = logic.parse_taught_log

    def spy(table: Table) -> tuple[list[logic.TaughtEntry], list[str]]:
        seen.append(table)
        return parse(table)

    monkeypatch.setattr(logic, "parse_taught_log", spy)
    records = _Records()
    csvio_log = logging.getLogger("staffroom.core.csvio")
    csvio_log.addHandler(records)
    logged: list[str] = []
    ctx = RunContext(config=load_config(), output_dir=tmp_path / "out", log=logged.append)
    options = Options(fmt="forms", retrieval=True, today="2026-10-08", count=3, seed=1)
    try:
        result = plugin.run(
            {"source": EXAMPLES / "question_bank.md", "taught_log": log_path, "options": options},
            ctx,
        )
    finally:
        csvio_log.removeHandler(records)

    # The parser only ever saw date and topic.
    assert len(seen) == 1
    assert seen[0].columns == ["date", "topic"]
    assert seen[0].dropped_columns == ["name", "Class"]
    assert not any("CELL" in str(c) for row in seen[0].rows for c in row)
    # The usual csvio warning, and the tool's own line, by header only.
    assert records.messages == ["taught.csv: dropped identifier column(s) name, Class"]
    assert "Log: dropped identifier column(s), never read: name, Class" in logged
    # The run still works from the surviving columns, and no value leaks anywhere.
    assert result.tables["chosen"].rows
    everywhere = [*logged, *records.messages, result.summary, *result.texts.values()]
    assert not any("CELL" in m for m in everywhere)


def test_taught_log_must_be_a_table_file(tmp_path: Path) -> None:
    log_path = tmp_path / "taught.txt"
    log_path.write_text("date,topic\n2026-10-01,networks\n", encoding="utf-8")
    ctx = RunContext(config=load_config(), output_dir=tmp_path / "out")
    options = Options(fmt="forms", retrieval=True, today="2026-10-08", count=3, seed=1)
    with pytest.raises(ToolError, match=r"taught\.txt is a \.txt file"):
        plugin.run(
            {"source": EXAMPLES / "question_bank.md", "taught_log": log_path, "options": options},
            ctx,
        )


def test_kahoot_rows_and_files(bank: logic.Bank, tmp_path: Path) -> None:
    sheet = logic.kahoot_sheet(bank.questions)
    rows = sheet.rows
    assert all(len(r) == 7 for r in rows)
    assert len(rows) == len(sheet.written) == 13 and sheet.skipped == ()
    assert rows[0][0] == "What does CPU stand for?" and rows[0][6] == "1" and rows[0][5] == "20"
    tf = next(r for r in rows if r[0] == "RAM is volatile.")
    assert tf[1:5] == ["True", "False", "", ""] and tf[6] == "1"
    short = next(r for r in rows if r[0] == "What is 1010 in denary?")
    assert "10" in short[1:5] and short[6] == str(short[1:5].index("10") + 1)
    smtp = next(r for r in rows if r[0].startswith("Name the protocol"))
    assert "SMTP" in smtp[1:5] and all(smtp[1:5])  # distractors borrowed from other short answers
    assert all(int(r[5]) in logic.KAHOOT_TIMES for r in rows)

    files = write_kahoot(rows, tmp_path, "bank")
    assert [f.name for f in files] == ["bank_kahoot.xlsx", "bank_kahoot.csv"]
    ws = load_workbook(files[0]).active
    assert ws is not None
    headers = [ws.cell(row=8, column=c).value for c in range(2, 9)]
    assert tuple(headers) == KAHOOT_HEADERS
    assert ws.cell(row=9, column=1).value == 1 and ws.cell(row=9, column=2).value == rows[0][0]
    assert ws.cell(row=9, column=7).value == 20
    reader = csv.reader(io.StringIO(files[1].read_text().split("\n", 1)[1]))
    assert tuple(next(reader)) == KAHOOT_HEADERS


def test_clipping_and_time_snapping() -> None:
    long = logic.Question(1, "x" * 130, "mcq", ("a" * 80, "b"), (0,), "", (), 25)
    sheet = logic.kahoot_sheet([long])
    rows = sheet.rows
    assert len(rows[0][0]) == 120 and len(rows[0][1]) == 75 and rows[0][5] == "20"
    assert len(sheet.notes) == 2


def _short(number: int, text: str, answer: str, tags: str) -> logic.Question:
    return logic.Question(number, text, "short", (), (), answer, tuple(tags.split(",")))


def test_kahoot_distractors_prefer_topic_then_spec_then_anywhere() -> None:
    # Q1 is the question; the only same-topic donor is Q2, the only same-spec-section
    # donor is Q3, and Q4..Q6 are unrelated. All three tiers must be used, nearest first.
    q = _short(1, "1010 in denary?", "10", "ocr:1.2.4,topic:binary")
    pool = [
        q,
        _short(2, "255 in binary?", "11111111", "ocr:1.2.4,topic:binary"),
        _short(3, "Bits in a byte?", "8", "ocr:1.2.3,topic:data"),
        _short(4, "Email protocol?", "SMTP", "ocr:1.3.2,topic:networks"),
        _short(5, "Web protocol?", "HTTP", "ocr:1.3.2,topic:networks"),
        _short(6, "Secure web protocol?", "HTTPS", "ocr:1.3.2,topic:networks"),
    ]
    sheet = logic.kahoot_sheet([q], pool)
    options = set(sheet.rows[0][1:5])
    assert {"10", "11111111", "8"} <= options and len(options & {"SMTP", "HTTP", "HTTPS"}) == 1
    # Within the topic there is nothing to say; outside it the note names the sources.
    assert len(sheet.notes) == 1
    note = sheet.notes[0]
    assert note.startswith("Q1: ") and "outside topic 'binary'" in note
    assert "'8' (Q3, data)" in note
    assert any(f"(Q{n}, networks)" in note for n in (4, 5, 6))

    # Enough same-topic donors: nothing borrowed from elsewhere, no note.
    binary = [
        q,
        _short(2, "a", "11111111", "topic:binary"),
        _short(3, "b", "1F", "topic:binary"),
        _short(4, "c", "8", "topic:binary"),
        _short(5, "d", "SMTP", "topic:networks"),
    ]
    sheet = logic.kahoot_sheet([q], binary)
    assert set(sheet.rows[0][1:5]) == {"10", "11111111", "1F", "8"} and sheet.notes == []


def test_kahoot_borrows_from_the_whole_bank_and_dedupes(bank: logic.Bank) -> None:
    # Two binary questions chosen out of the bank: on their own they have one distractor
    # each, so the bank's other short answers must supply the rest, same topic first.
    chosen = [q for q in bank.questions if q.number in (5, 6)]
    sheet = logic.kahoot_sheet(chosen, bank.questions)
    assert [q.number for q in sheet.written] == [5, 6] and sheet.skipped == ()
    for row in sheet.rows:
        assert set(row[1:5]) <= {"10", "11111111", "1F", "8"}  # never SMTP
    assert sheet.notes == []
    # Answers equal to the question's own, or repeated among donors, count once.
    q = _short(1, "x", "8", "topic:t")
    pool = [q, _short(2, "a", "8", "topic:t"), _short(3, "b", " 8 ", "topic:t")]
    pool += [_short(4, "c", "9", "topic:t"), _short(5, "d", "10", "topic:t")]
    sheet = logic.kahoot_sheet([q], pool)
    assert sheet.rows == [] and len(sheet.skipped) == 1
    assert "too few distractors" in sheet.skipped[0][1] and "found 2" in sheet.skipped[0][1]


def test_kahoot_sheet_reports_what_was_written(bank: logic.Bank) -> None:
    lonely = _short(14, "Which gate outputs 1 only when both inputs are 1?", "AND", "topic:logic")
    chosen = [bank.questions[0], bank.questions[2], lonely]
    sheet = logic.kahoot_sheet(chosen, [*bank.questions[:4], lonely])  # no other short answers
    assert [q.number for q in sheet.written] == [1, 3] and len(sheet.rows) == 2
    assert [q.number for q, _ in sheet.skipped] == [14]
    assert logic.kahoot_report(sheet, 3) == "2 of 3 written; 1 short answer skipped: Q14"
    assert logic.kahoot_report(logic.kahoot_sheet(bank.questions[:2]), 2) == "2 questions"
    # An mcq whose correct option is beyond the fourth is skipped too, and the summary says so
    # without calling it a short answer.
    five = logic.Question(20, "Pick", "mcq", ("a", "b", "c", "d", "e"), (4,), "", ())
    sheet = logic.kahoot_sheet([bank.questions[0], five, lonely], [*bank.questions[:4], lonely])
    assert logic.kahoot_report(sheet, 3) == "1 of 3 written; 2 skipped: Q20, Q14"


def test_plugin_kahoot_summary_and_table_match_the_file(tmp_path: Path) -> None:
    # The first four example questions (two mcq, two tf: no short answers at all) plus
    # one short answer that therefore has nothing to borrow from.
    text = (EXAMPLES / "question_bank.md").read_text().split("## What is 1010 in denary?")[0]
    text += "## Which gate outputs 1 only when both inputs are 1?\ntype: short\n"
    text += "tags: topic:logic\nanswer: AND\n"
    small = tmp_path / "small.md"
    small.write_text(text, encoding="utf-8")
    out = tmp_path / "out"
    logged: list[str] = []
    ctx = RunContext(config=load_config(), output_dir=out, log=logged.append)
    options = Options(fmt="kahoot", count=5, seed=1)  # picks all five
    result = plugin.run({"source": small, "options": options}, ctx)
    numbers = [row[0] for row in result.tables["chosen"].rows]
    assert numbers == [1, 2, 3, 4], "the skipped short answer must not be in the chosen table"
    body = (out / "small_kahoot.csv").read_text().split("\n", 1)[1]  # after the FAKE marker
    data_rows = list(csv.reader(io.StringIO(body)))[1:]
    assert len(data_rows) == len(numbers) == 4
    assert result.summary.startswith("4 of 5 written; 1 short answer skipped: Q5 → small_kahoot")
    assert any("Q5 skipped for Kahoot" in m and "found 0" in m for m in logged)

    # When nothing can be written the teacher gets a ToolError, not an empty file.
    only = tmp_path / "only.md"
    only.write_text("## Gate?\ntype: short\ntags: topic:logic\nanswer: AND\n", encoding="utf-8")
    with pytest.raises(ToolError, match="Nothing could be exported for Kahoot"):
        plugin.run({"source": only, "options": Options(fmt="kahoot", count=1)}, ctx)


def test_forms_text_and_donow(bank: logic.Bank, tmp_path: Path) -> None:
    text = logic.forms_text(bank.questions[:3], "Quick quiz")
    assert text.startswith(
        "Quick quiz\n\n1. What does CPU stand for?\n   A) Central Processing Unit"
    )
    assert "Answers\n\n1. A\n2. B\n3. A\n" in text
    ctx = RunContext(config=load_config(), output_dir=tmp_path)
    result = plugin.run(
        {
            "source": EXAMPLES / "question_bank.md",
            "options": Options(fmt="donow", count=4, seed=2, title="Do now"),
        },
        ctx,
    )
    assert sorted(p.name for p in tmp_path.iterdir()) == [
        "question_bank_donow.docx",
        "question_bank_donow.pdf",
        "question_bank_donow_key.docx",
        "question_bank_donow_key.pdf",
    ]
    assert "ANSWER KEY" in result.texts["preview"] and result.tables["chosen"].shape[0] == 4


def test_plugin_retrieval_and_cli(tmp_path: Path) -> None:
    ctx = RunContext(config=load_config(), output_dir=tmp_path)
    options = Options(fmt="forms", retrieval=True, today="2026-10-08", count=6, seed=1)
    result = plugin.run(
        {
            "source": EXAMPLES / "question_bank.md",
            "taught_log": EXAMPLES / "taught_log.csv",
            "options": options,
        },
        ctx,
    )
    topics = {str(row[3]) for row in result.tables["chosen"].rows}
    assert topics <= {"networks", "loops", "cpu"}
    r = runner.invoke(
        app,
        [
            "quiz",
            "export",
            "--bank",
            str(EXAMPLES / "question_bank.md"),
            "-f",
            "kahoot",
            "-t",
            "binary",
            "-n",
            "3",
            "--out",
            str(tmp_path),
        ],
    )
    assert r.exit_code == 0, r.output
    assert (tmp_path / "question_bank_kahoot.xlsx").exists()
    r = runner.invoke(app, ["quiz", "list", "--bank", str(EXAMPLES / "question_bank.md")])
    assert r.exit_code == 0 and "13 questions" in r.output
    r = runner.invoke(app, ["quiz", "check", "--bank", str(EXAMPLES / "question_bank.md")])
    assert r.exit_code == 0 and "0 problems" in r.output and "0 unknown spec tags" in r.output
    r = runner.invoke(app, ["quiz", "spec", "aqa"])
    assert r.exit_code == 0 and "aqa:3.3.4" in r.output
    r = runner.invoke(
        app,
        [
            "quiz",
            "retrieval",
            "--bank",
            str(EXAMPLES / "question_bank.md"),
            "--log",
            str(EXAMPLES / "taught_log.csv"),
            "--date",
            "2026-10-08",
            "--out",
            str(tmp_path),
        ],
    )
    assert r.exit_code == 0, r.output
    assert "networks" in r.output or "loops" in r.output
