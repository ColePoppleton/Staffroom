"""Acceptance: every variant differs from the original by exactly the
described mutation, and the key says what each one does."""

from collections.abc import Callable
from pathlib import Path

import pytest
from pypdf import PdfReader
from typer.testing import CliRunner

from staffroom.cli import app
from staffroom.core import paths
from staffroom.core.config import load_config
from staffroom.core.export import render, to_text
from staffroom.core.plugin import RunContext, RunResult, ToolError
from staffroom.core.runner import RunOutcome, run_source
from staffroom.tools.bugs import classify, logic, outcome_kind, plugin, runs_the_same
from staffroom.tools.bugs.logic import KIND_LABELS, Options
from staffroom.tools.bugs.sheet import build_document

runner = CliRunner()
YEAR9 = paths.examples_dir() / "year9.py"

PROGRAM = """total = 0
count = 1
limit = 5
while count <= limit:
    total = total + count
    count = count + 1
for i in range(3):
    print(i)
score = int(input())
if score > 10:
    print("high")
else:
    print("low")
print(total)
"""


def test_candidates_cover_every_kind() -> None:
    kinds = {m.kind for m in logic.candidates(PROGRAM)}
    assert kinds == set(KIND_LABELS)


@pytest.mark.parametrize("seed", range(1, 12))
def test_each_variant_is_exactly_one_described_mutation(seed: int) -> None:
    options = Options(count=7, seed=seed)
    variants = logic.make_variants(PROGRAM, options)
    assert len(variants) == 7
    assert len({v.mutation.kind for v in variants}) == 7  # one of each kind first
    for v in variants:
        diff = logic.one_line_diff(PROGRAM, v.source)
        assert diff is not None, v.mutation
        line_no, old, new = diff
        assert (
            line_no == v.mutation.line and old == v.mutation.original and new == v.mutation.mutated
        )
        assert logic.describes(v.mutation, old, new), (v.mutation.kind, old, new)
    again = logic.make_variants(PROGRAM, options)
    assert [v.source for v in again] == [v.source for v in variants]


def test_kinds_can_be_restricted() -> None:
    options = Options(
        count=3,
        seed=2,
        comparison_flip=False,
        off_by_one=False,
        missing_colon=False,
        missing_bracket=False,
        wrong_variable=False,
        indentation_slip=False,
    )
    variants = logic.make_variants(PROGRAM, options)
    assert variants and all(v.mutation.kind == "string_int" for v in variants)


def test_plugin_runs_variants_and_classifies(tmp_path: Path) -> None:
    logs: list[str] = []
    ctx = RunContext(config=load_config(), output_dir=tmp_path, log=logs.append)
    result = plugin.run(
        {"pasted": PROGRAM, "stdin": "7\n", "options": Options(count=7, seed=3, formats="pdf")}, ctx
    )
    names = sorted(p.name for p in tmp_path.iterdir())
    assert names[:7] == [f"pasted_code_bug{i}.py" for i in range(1, 8)]
    assert "pasted_code_bugs.pdf" in names and "pasted_code_bugs_key.pdf" in names
    outcomes = [str(row[5]) for row in result.tables["key"].rows]
    assert any(o.startswith("crashes") for o in outcomes)
    assert any("output is wrong" in o or "no visible change" in o for o in outcomes)
    for row in result.tables["key"].rows:  # each file is the key's one-line change
        diff = logic.one_line_diff(PROGRAM, (tmp_path / f"pasted_code_bug{row[0]}.py").read_text())
        assert diff is not None
        line_no, old, new = diff
        shown = logic.key_lines(logic.Mutation("off_by_one", line_no, old, new, ""))
        assert (line_no, *shown) == (row[1], row[3], row[4])


def test_cli(tmp_path: Path) -> None:
    src = tmp_path / "prog.py"
    src.write_text(PROGRAM)
    r = runner.invoke(
        app,
        [
            "bugs",
            "make",
            str(src),
            "-n",
            "3",
            "--stdin",
            "7",
            "--kinds",
            "missing_colon,off_by_one",
            "-f",
            "docx",
            "--out",
            str(tmp_path),
        ],
    )
    assert r.exit_code == 0, r.output
    assert (
        "line" in r.output
        and (tmp_path / "prog_bug1.py").exists()
        and (tmp_path / "prog_bugs_key.docx").exists()
    )
    r = runner.invoke(app, ["bugs", "kinds"])
    assert r.exit_code == 0 and "indentation_slip" in r.output
    r = runner.invoke(app, ["bugs", "make", str(src), "--kinds", "nope"])
    assert r.exit_code == 1


def test_classify_says_when_a_variant_prints_past_the_cap() -> None:
    original = RunOutcome(stdout="1\n", returncode=0)
    flood = RunOutcome(stdout="1\n1\n", returncode=0, stdout_truncated=True)
    assert classify(original, flood) == "runs, but the output is wrong; output capped at 1024 KB"
    assert classify(original, RunOutcome(stdout="1\n", returncode=0)).startswith("silent")


def _pages(doc_variants: int, tmp_path: Path) -> int:
    source = YEAR9.read_text()
    variants = logic.make_variants(source, Options(count=doc_variants, seed=1))
    assert len(variants) == doc_variants
    doc = build_document("year9", source, variants)
    path = render(doc.sheet(), tmp_path / f"bugs{doc_variants}.pdf")
    return len(PdfReader(str(path)).pages)


def test_two_variants_per_printed_page(tmp_path: Path) -> None:
    """The audit's measured failure (#22): four 13-line variants printed on
    four pages. Two per page, with the answer line under each."""
    assert _pages(4, tmp_path) == 2
    assert _pages(1, tmp_path) == 1
    assert _pages(5, tmp_path) == 3


def test_page_break_only_between_pairs() -> None:
    source = YEAR9.read_text()
    variants = logic.make_variants(source, Options(count=4, seed=1))
    text = to_text(build_document("year9", source, variants))
    assert text.count("page break") == 1
    v1 = text.index("Version 1")
    v2 = text.index("Version 2")
    assert text.index("Line: ______", v1) < v2  # the answer line sits under its version
    assert text.index("page break") > text.index("Line: ______", v2)


# ---------------------------------------------------------------- #92: bugs a pupil can see
def _run(
    tmp_path: Path, options: Options, source: Path | None = None
) -> tuple[RunResult, list[str], list[str]]:
    logs: list[str] = []
    warnings: list[str] = []
    ctx = RunContext(
        config=load_config(), output_dir=tmp_path, log=logs.append, warned=warnings.append
    )
    result = plugin.run({"source": source or YEAR9, "options": options}, ctx)
    return result, logs, warnings


def test_default_year9_variants_all_change_behaviour(tmp_path: Path) -> None:
    """The audit's measured failure: at the defaults two of four year9 variants
    ran exactly like the original. Now all four change what a pupil sees."""
    result, logs, warnings = _run(tmp_path, Options(formats="pdf"))
    rows = result.tables["key"].rows
    assert len(rows) == 4
    assert not any("silent" in str(row[5]) for row in rows)
    original = run_source(YEAR9.read_text())
    for row in rows:
        variant = run_source((tmp_path / f"year9_bug{row[0]}.py").read_text())
        assert not runs_the_same(original, variant), row
    assert not warnings
    assert any(line.startswith("Skipped line 9") for line in logs)  # the two silent ones
    sheet = result.texts["preview"]
    assert "read the code" not in sheet and "(crash or wrong output)" in sheet
    assert "seed 1" in result.outcome


def test_same_seed_same_variants(tmp_path: Path) -> None:
    first, _, _ = _run(tmp_path / "a", Options(seed=5, formats="pdf"))
    second, _, _ = _run(tmp_path / "b", Options(seed=5, formats="pdf"))
    assert first.tables["key"].rows == second.tables["key"].rows


def test_kept_silent_variant_is_labelled(tmp_path: Path) -> None:
    result, _, warnings = _run(tmp_path, Options(silent=1, formats="pdf"))
    rows = result.tables["key"].rows
    silent = [row for row in rows if str(row[5]).startswith("silent")]
    assert len(rows) == 4 and len(silent) == 1
    sheet = result.texts["preview"]
    assert f"Version {silent[0][0]} (this version may run the same: read the code)" in sheet
    assert sheet.count("read the code") == 1
    assert "1 silent" in result.outcome
    assert any("runs exactly like the original" in w for w in warnings)


def test_says_when_fewer_visible_bugs_are_possible(tmp_path: Path) -> None:
    result, _, warnings = _run(tmp_path, Options(count=20, formats="pdf"))
    rows = result.tables["key"].rows
    assert 0 < len(rows) < 20
    assert not any("silent" in str(row[5]) for row in rows)
    assert any(f"Only {len(rows)} distinct bugs change" in w for w in warnings), warnings
    assert "(asked for 20)" in result.outcome


def test_every_bug_silent_is_a_teacher_error(tmp_path: Path) -> None:
    src = tmp_path / "quiet.py"
    src.write_text('x = 1\nif x > 5:\n    print("big")\n')
    only_flips = Options(
        formats="pdf",
        off_by_one=False,
        missing_colon=False,
        missing_bracket=False,
        wrong_variable=False,
        indentation_slip=False,
        string_int=False,
    )
    with pytest.raises(ToolError, match="Silent bugs allowed"):
        _run(tmp_path / "out", only_flips, src)
    only_flips.silent = 1
    result, _, _ = _run(tmp_path / "out", only_flips, src)
    assert str(result.tables["key"].rows[0][5]).startswith("silent")


def test_choose_skips_silent_candidates_in_seed_order() -> None:
    every = logic.choose(PROGRAM, 30, list(KIND_LABELS), seed=4)
    quiet = {m for m in every if m.kind in ("comparison_flip", "off_by_one")}
    asked: list[logic.Mutation] = []

    def changes(m: logic.Mutation) -> bool:
        asked.append(m)
        return m not in quiet

    picked = logic.choose(PROGRAM, 4, list(KIND_LABELS), 4, changes)
    assert len(picked) == 4 and not quiet & set(picked)
    assert len(asked) == len(set(asked))  # each candidate is run at most once
    assert logic.choose(PROGRAM, 4, list(KIND_LABELS), 4, changes) == picked  # reproducible
    kept_one = logic.choose(PROGRAM, 4, list(KIND_LABELS), 4, changes, silent=1)
    assert len(quiet & set(kept_one)) == 1
    # Allowing as many silent ones as variants is the old, unchecked choice.
    assert logic.choose(PROGRAM, 4, list(KIND_LABELS), 4, changes, silent=4) == logic.choose(
        PROGRAM, 4, list(KIND_LABELS), 4
    )


def test_same_crash_as_a_crashing_original_is_silent() -> None:
    boom = RunOutcome(stdout="a\n", stderr="ZeroDivisionError: division by zero", returncode=1)
    assert runs_the_same(boom, boom)
    assert classify(boom, boom).startswith("silent")
    other = RunOutcome(stdout="a\n", stderr="NameError: name 'x' is not defined", returncode=1)
    assert classify(boom, other) == "crashes: NameError"


# ---------------------------------------------------------------- #118: a mix of outcomes
_BY_KIND: dict[str, logic.Outcome] = {
    "missing_colon": "syntax_error",
    "missing_bracket": "syntax_error",
    "indentation_slip": "hang",
}


def _fake_outcomes(wrong: logic.Mutation | None) -> Callable[[logic.Mutation], logic.Outcome]:
    """Every bug crashes (colons and brackets before the program starts),
    except ``wrong``, the one bug that runs with the wrong output."""

    def outcome_of(m: logic.Mutation) -> logic.Outcome:
        if m == wrong:
            return "wrong_output"
        return _BY_KIND.get(m.kind, "crash")

    return outcome_of


def test_choose_swaps_in_a_wrong_output_bug_and_holds_back_syntax_errors() -> None:
    every = logic.choose(PROGRAM, 100, list(KIND_LABELS), seed=1)
    wrong = every[-1]  # the last candidate the seed's order reaches
    picked = logic.choose(PROGRAM, 4, list(KIND_LABELS), 1, _fake_outcomes(wrong))
    outcomes = [_fake_outcomes(wrong)(m) for m in picked]
    assert len(picked) == 4 and wrong in picked
    assert outcomes.count("syntax_error") <= 1
    assert logic.choose(PROGRAM, 4, list(KIND_LABELS), 1, _fake_outcomes(wrong)) == picked
    # One variant has no mix to aim for: the seed's first pick stands.
    one = logic.choose(PROGRAM, 1, list(KIND_LABELS), 1, _fake_outcomes(wrong))
    assert one == logic.choose(PROGRAM, 1, list(KIND_LABELS), 1)


def test_choose_without_the_mix_still_makes_up_the_count() -> None:
    """No wrong-output bug anywhere, and only SyntaxErrors in the ticked
    classes: the count is still met, with as many SyntaxErrors as it takes."""
    syntax_only = ["missing_colon", "missing_bracket"]
    picked = logic.choose(PROGRAM, 3, syntax_only, 2, _fake_outcomes(None))  # type: ignore[arg-type]
    assert len(picked) == 3
    assert {m.kind for m in picked} <= set(syntax_only)
    mixed = logic.choose(PROGRAM, 4, list(KIND_LABELS), 2, _fake_outcomes(None))
    assert len(mixed) == 4
    assert [_fake_outcomes(None)(m) for m in mixed].count("syntax_error") <= 1


def test_choose_never_swaps_out_a_kept_silent_bug() -> None:
    """Silent bugs the teacher allowed are kept, not traded for the mix."""
    wrong = logic.choose(PROGRAM, 100, list(KIND_LABELS), seed=1)[-1]

    def silent_but_one(m: logic.Mutation) -> logic.Outcome:
        return "wrong_output" if m == wrong else "silent"

    picked = logic.choose(PROGRAM, 3, list(KIND_LABELS), 1, silent_but_one, silent=3)
    assert picked == logic.choose(PROGRAM, 3, list(KIND_LABELS), 1)


@pytest.fixture(scope="module")
def year9_outcomes() -> dict[logic.Mutation, logic.Outcome]:
    """What every year9.py candidate really does, run once in the sandbox."""
    source = YEAR9.read_text()
    original = run_source(source, timeout=logic.TIMEOUT)
    return {
        m: outcome_kind(original, run_source(logic.apply(source, m), timeout=logic.TIMEOUT))
        for m in logic.candidates(source)
    }


@pytest.mark.parametrize("seed", range(1, 11))
def test_year9_default_sheet_mixes_outcomes(
    seed: int, year9_outcomes: dict[logic.Mutation, logic.Outcome]
) -> None:
    source = YEAR9.read_text()
    picked = logic.choose(source, 4, list(KIND_LABELS), seed, year9_outcomes.__getitem__)
    outcomes = [year9_outcomes[m] for m in picked]
    assert len(picked) == 4
    assert "wrong_output" in outcomes and "silent" not in outcomes
    assert outcomes.count("syntax_error") <= 1


def test_default_year9_sheet_has_a_wrong_output_bug(tmp_path: Path) -> None:
    """The audit's failure (#118): at seed 1 all four versions crashed, two
    of them with SyntaxErrors. Now one runs with the wrong output."""
    result, logs, warnings = _run(tmp_path, Options(formats="pdf"))
    outcomes = [str(row[5]) for row in result.tables["key"].rows]
    assert len(outcomes) == 4
    assert "runs, but the output is wrong" in outcomes
    assert sum("SyntaxError" in o for o in outcomes) == 1
    assert not warnings
    assert any("already has a SyntaxError" in line for line in logs)


def test_says_when_the_program_cannot_offer_the_mix(tmp_path: Path) -> None:
    syntax_only = Options(
        count=2,
        formats="pdf",
        comparison_flip=False,
        off_by_one=False,
        wrong_variable=False,
        indentation_slip=False,
        string_int=False,
    )
    result, _, warnings = _run(tmp_path, syntax_only)
    outcomes = [str(row[5]) for row in result.tables["key"].rows]
    assert outcomes == ["crashes: SyntaxError"] * 2
    assert any("None of the possible bugs" in w and "Ticking more" in w for w in warnings)
    assert any(w.startswith("2 of the 2 versions stop with a SyntaxError") for w in warnings)


def test_outcome_kind_names_compile_errors_and_hangs() -> None:
    ok = RunOutcome(stdout="1\n", returncode=0)
    assert outcome_kind(ok, ok) == "silent"
    assert outcome_kind(ok, RunOutcome(stdout="2\n", returncode=0)) == "wrong_output"
    assert outcome_kind(ok, RunOutcome(stdout="", timed_out=True)) == "hang"
    for name in ("SyntaxError", "IndentationError"):
        broken = RunOutcome(stderr=f"{name}: invalid syntax", returncode=1)
        assert outcome_kind(ok, broken) == "syntax_error"
    assert outcome_kind(ok, RunOutcome(stderr="TypeError: no", returncode=1)) == "crash"


# ---------------------------------------------------------------- #127: a plausible wrong name
def _swaps(source: str) -> dict[tuple[int, str], str]:
    """(line, original name) -> the name the wrong_variable bug puts in."""
    found = {}
    for m in logic.candidates(source):
        if m.kind == "wrong_variable":
            gone, came = m.note.split(" replaced by ") if "replaced" in m.note else (None, None)
            if gone and came:
                found[(m.line, gone)] = came
    return found


def test_wrong_variable_uses_a_name_of_the_same_kind() -> None:
    """The audit's failure (#127): every swap in year9.py brought in the
    list ``scores``, so ``if scores >= 50:`` crashed. Now a number stands in
    for a number, the one most recently assigned before the line."""
    swaps = _swaps(YEAR9.read_text())
    assert swaps[(9, "average")] == "count"  # if count >= 50:
    assert swaps[(6, "total")] == "score"  # total = score + score
    assert swaps[(6, "score")] == "count"
    assert swaps[(8, "total")] == "count" and swaps[(8, "count")] == "total"
    numbers = {"total", "count", "average", "score"}
    assert all(came in numbers for (_, gone), came in swaps.items() if gone in numbers)
    assert logic.candidates(YEAR9.read_text()) == logic.candidates(YEAR9.read_text())


def test_wrong_variable_stand_in_is_in_scope_and_already_assigned() -> None:
    source = (
        'name = input("Name? ")\n'
        'greeting = "Hello"\n'
        "def area(width, height):\n"
        "    side = width\n"
        "    return side * height\n"
        "size = area(2, 3)\n"
        "later = 4\n"
        "print(greeting, name, size, later)\n"
    )
    swaps = _swaps(source)
    assert swaps[(8, "greeting")] == "name" and swaps[(8, "name")] == "greeting"  # strings
    assert swaps[(5, "side")] in {"width", "height"}  # the function's own names
    assert swaps[(8, "size")] == "later"  # a number, assigned before line 8
    assert all(
        came not in {"width", "height", "side"} for (line, _), came in swaps.items() if line > 5
    )


def test_year9_wrong_variable_bugs_mostly_change_the_output(
    year9_outcomes: dict[logic.Mutation, logic.Outcome],
) -> None:
    """Before #127 seven of the eight year9 wrong-name bugs were TypeErrors."""
    wrong_names = {m: o for m, o in year9_outcomes.items() if m.kind == "wrong_variable"}
    assert len(wrong_names) >= 6
    crashes = [m for m, o in wrong_names.items() if o == "crash"]
    assert len(crashes) <= 1, crashes  # only `for score in score:`, a NameError
    assert sum(o == "wrong_output" for o in wrong_names.values()) >= len(wrong_names) - 1


# ---------------------------------------------------------------- #130: indentation in the key
def test_key_lines_show_indentation_only_when_it_changed() -> None:
    slip = logic.Mutation("indentation_slip", 7, "    count = count + 1", "count = count + 1", "")
    assert logic.key_lines(slip) == ("····count = count + 1", "count = count + 1")
    tab = logic.Mutation("indentation_slip", 2, "\t\tx = 1", "\tx = 1", "")
    assert logic.key_lines(tab) == ("»»x = 1", "»x = 1")
    flip = logic.Mutation("comparison_flip", 9, "    if a > 1:", "    if a >= 1:", "")
    assert logic.key_lines(flip) == ("if a > 1:", "if a >= 1:")
    assert logic.key_label(slip) == "indentation slip (4 spaces to 0)"
    assert logic.key_label(flip) == "comparison flipped"


def test_year9_seed1_key_shows_the_indentation_slip(tmp_path: Path) -> None:
    """The audit's failure (#130): at seed 1 the indentation bug's key row read
    ``count = count + 1`` twice. Every key now shows the indentation."""
    from docx import Document as Docx

    result, _, _ = _run(tmp_path, Options(seed=1, formats="both"))
    slips = [row for row in result.tables["key"].rows if "indentation" in str(row[2])]
    assert len(slips) == 1
    _, line, bug, old, new, _ = slips[0]
    assert bug == "indentation slip (4 spaces to 0)"
    assert line == 7 and (old, new) == ("····count = count + 1", "count = count + 1")
    assert old != new and str(old).replace("·", "") == new  # differ only in the visible indent

    text = result.texts["preview"]
    assert "····count = count + 1" in text and logic.INDENT_LEGEND in text
    pdf = " ".join(p.extract_text() for p in PdfReader(str(tmp_path / "year9_bugs_key.pdf")).pages)
    pdf = " ".join(pdf.split())  # the narrow PDF cells wrap
    assert "····count = count + 1 count = count + 1" in pdf and logic.INDENT_LEGEND in pdf
    assert "(4 spaces to 0)" in pdf
    cells = [c.text for t in Docx(str(tmp_path / "year9_bugs_key.docx")).tables for c in t._cells]
    assert "····count = count + 1" in cells and "count = count + 1" in cells
    assert "indentation slip (4 spaces to 0)" in cells


def test_key_without_an_indentation_bug_has_no_legend() -> None:
    source = YEAR9.read_text()
    variants = logic.make_variants(source, Options(count=2, seed=1, indentation_slip=False))
    text = to_text(build_document("year9", source, variants))
    assert "·" not in text.split("ANSWER KEY")[1] and logic.INDENT_LEGEND not in text
