"""Read and write tables as CSV or XLSX.

Fake-data marker
----------------
Real pupil data must never be committed. Every example file carries a marker
that the pre-commit hook looks for:

* CSV: the first line is ``# FAKE DATA`` (anything after the words is ignored).
* XLSX: some cell in the first few rows of any sheet contains ``FAKE DATA``.

:func:`read_table` strips the marker and records it on ``Table.fake_data``;
:func:`write_table` puts it back, so derived outputs stay committable.
"""

from __future__ import annotations

import csv
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook, load_workbook

from teaching_tools.core.table import Row, Table, coerce

FAKE_DATA_MARKER = "# FAKE DATA"
_MARKER_WORDS = "FAKE DATA"
_MARKER_SCAN_ROWS = 3

TABLE_EXTENSIONS = (".csv", ".xlsx")


def is_marker_line(line: str) -> bool:
    return line.strip().upper().startswith(FAKE_DATA_MARKER)


def read_table(path: Path) -> Table:
    """Read a CSV or XLSX file, dispatching on suffix."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    msg = f"unsupported table file {path.name!r}; expected one of {TABLE_EXTENSIONS}"
    raise ValueError(msg)


def write_table(table: Table, path: Path) -> Path:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return write_csv(table, path)
    if suffix == ".xlsx":
        return write_xlsx(table, path)
    msg = f"unsupported table file {path.name!r}; expected one of {TABLE_EXTENSIONS}"
    raise ValueError(msg)


# ---------------------------------------------------------------- CSV
def read_csv(path: Path) -> Table:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        lines = fh.read().splitlines()
    fake = False
    while lines and (is_marker_line(lines[0]) or not lines[0].strip()):
        fake = fake or is_marker_line(lines[0])
        lines.pop(0)
    if not lines:
        return Table([], [], fake_data=fake)
    reader = csv.reader(lines)
    header = [h.strip() for h in next(reader)]
    rows: list[Row] = [[coerce(c) for c in _pad(r, len(header))] for r in reader if any(r)]
    return Table(header, rows, fake_data=fake)


def write_csv(table: Table, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        if table.fake_data:
            fh.write(FAKE_DATA_MARKER + "\n")
        writer = csv.writer(fh)
        writer.writerow(table.columns)
        for row in table.rows:
            writer.writerow(["" if c is None else c for c in row])
    return path


# ---------------------------------------------------------------- XLSX
def read_xlsx(path: Path, sheet: str | None = None) -> Table:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    raw = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    fake = xlsx_has_marker(path)
    # Drop marker/blank rows before the header.
    while raw and (_row_is_marker(raw[0]) or not any(c is not None for c in raw[0])):
        raw.pop(0)
    if not raw:
        return Table([], [], fake_data=fake)
    header = ["" if c is None else str(c).strip() for c in raw[0]]
    rows: list[Row] = [
        [coerce(c) for c in _pad(r, len(header))] for r in raw[1:] if any(c is not None for c in r)
    ]
    return Table(header, rows, fake_data=fake)


def write_xlsx(table: Table, path: Path, sheet: str = "Sheet1") -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet
    if table.fake_data:
        ws.append([FAKE_DATA_MARKER])
    ws.append(table.columns)
    for row in table.rows:
        ws.append(row)
    wb.save(path)
    return path


def xlsx_has_marker(path: Path) -> bool:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _MARKER_SCAN_ROWS:
                    break
                if _row_is_marker(list(row)):
                    return True
    finally:
        wb.close()
    return False


def csv_has_marker(path: Path) -> bool:
    with path.open(encoding="utf-8-sig", errors="replace") as fh:
        for _ in range(_MARKER_SCAN_ROWS):
            line = fh.readline()
            if is_marker_line(line):
                return True
    return False


def has_fake_data_marker(path: Path) -> bool:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return csv_has_marker(path)
    if suffix == ".xlsx":
        return xlsx_has_marker(path)
    return False


# ---------------------------------------------------------------- helpers
def _pad(row: Sequence[object], width: int) -> list[object]:
    padded = list(row[:width])
    padded.extend([None] * (width - len(padded)))
    return padded


def _row_is_marker(row: Sequence[object]) -> bool:
    return any(isinstance(c, str) and _MARKER_WORDS in c.upper() for c in row)
