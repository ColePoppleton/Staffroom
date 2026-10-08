"""docx renderer (python-docx) in the house style."""

from __future__ import annotations

from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph

from teaching_tools.core.export.model import (
    Code,
    Document,
    Heading,
    HintBox,
    Numbered,
    PageBreak,
    Paragraph,
    TableBlock,
)

BODY_FONT = "Calibri"
MONO_FONT = "Menlo"
MUTED = RGBColor(0x66, 0x70, 0x85)
ACCENT = RGBColor(0x3B, 0x5B, 0xDB)


def write_docx(doc: Document, path: Path) -> Path:
    d = docx.Document()
    styles = d.styles
    styles["Normal"].font.name = BODY_FONT
    styles["Normal"].font.size = Pt(11)
    _header_footer(d, doc)

    title = d.add_paragraph()
    run = title.add_run(doc.title)
    run.bold = True
    run.font.size = Pt(20)
    meta = "   ·   ".join(x for x in (doc.topic, doc.date_text, doc.subtitle) if x)
    if meta:
        p = d.add_paragraph()
        r = p.add_run(meta)
        r.font.color.rgb = MUTED

    for section in doc.sections:
        if section.title:
            d.add_heading(section.title, level=1)
        for block in section.blocks:
            _render(d, block)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path


def _header_footer(d: DocxDocument, doc: Document) -> None:
    section = d.sections[0]
    header = section.header.paragraphs[0]
    header.text = doc.header_text
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in header.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if doc.footer_text:
        lead = footer.add_run(doc.footer_text + "   ·   ")
        lead.font.size = Pt(9)
        lead.font.color.rgb = MUTED
    run = footer.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    _field(footer, "PAGE")
    run2 = footer.add_run(" of ")
    run2.font.size = Pt(9)
    run2.font.color.rgb = MUTED
    _field(footer, "NUMPAGES")


def _field(paragraph: DocxParagraph, code: str) -> None:
    """A Word field (PAGE / NUMPAGES), which Word fills in on open."""
    run = paragraph.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _mono(paragraph: DocxParagraph, text: str, size: int = 10) -> None:
    run = paragraph.add_run(text)
    run.font.name = MONO_FONT
    run.font.size = Pt(size)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), MONO_FONT)


def _render(d: DocxDocument, block: object) -> None:
    if isinstance(block, Heading):
        d.add_heading(block.text, level=min(max(block.level, 1), 4))
    elif isinstance(block, Paragraph):
        p = d.add_paragraph()
        r = p.add_run(block.text)
        if block.style == "muted":
            r.font.color.rgb = MUTED
        elif block.style == "instruction":
            r.italic = True
    elif isinstance(block, Code):
        if block.caption:
            cap = d.add_paragraph()
            cr = cap.add_run(block.caption)
            cr.font.color.rgb = MUTED
            cr.font.size = Pt(9)
        for line in block.text.splitlines() or [""]:
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Pt(12)
            _mono(p, line)
    elif isinstance(block, Numbered):
        for i, item in enumerate(block.items, block.start):
            first, *rest = item.splitlines() or [""]
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            num = p.add_run(f"{i}.  ")
            num.bold = True
            if block.code:
                _mono(p, first)
            else:
                p.add_run(first)
            for line in rest:
                q = d.add_paragraph()
                q.paragraph_format.left_indent = Pt(20)
                q.paragraph_format.space_after = Pt(0)
                if block.code:
                    _mono(q, line)
                else:
                    q.add_run(line)
            for _ in range(block.answer_lines):
                blank = d.add_paragraph()
                blank.paragraph_format.left_indent = Pt(20)
                br = blank.add_run("_" * 70)
                br.font.color.rgb = RGBColor(0xC9, 0xCE, 0xD9)
    elif isinstance(block, TableBlock):
        table = d.add_table(rows=1, cols=len(block.columns))
        table.style = "Light Grid Accent 1"
        for i, col in enumerate(block.columns):
            cell = table.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(col)
            run.bold = True
        for row in block.rows:
            cells = table.add_row().cells
            for i, value in enumerate(row):
                cells[i].text = ""
                if block.monospace:
                    _mono(cells[i].paragraphs[0], value)
                else:
                    cells[i].paragraphs[0].add_run(value)
        d.add_paragraph()
    elif isinstance(block, HintBox):
        table = d.add_table(rows=1, cols=1)
        table.style = "Light Shading Accent 1"
        cell = table.rows[0].cells[0]
        cell.text = ""
        p = cell.paragraphs[0]
        t = p.add_run(f"{block.title}: ")
        t.bold = True
        t.font.color.rgb = ACCENT
        p.add_run(block.text)
        d.add_paragraph()
    elif isinstance(block, PageBreak):
        d.add_page_break()  # type: ignore[no-untyped-call]
