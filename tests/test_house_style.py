from datetime import date
from pathlib import Path

from docx import Document as DocxReader
from pypdf import PdfReader

from teaching_tools.core.config import load_config
from teaching_tools.core.export import (
    Document,
    HouseStyle,
    Paragraph,
    export,
    house_style_from,
    to_text,
)


def test_house_style_appears_in_every_renderer(tmp_path: Path) -> None:
    doc = Document(title="Sheet", topic="Python", date=date(2026, 10, 8))
    doc.section().add(Paragraph("hello"))
    style = HouseStyle(header="Example High · Computing", footer="Computing Department")
    files = export(doc, tmp_path, "s", style=style)
    assert to_text(doc).splitlines()[2].startswith("Example High · Computing   ·   Python")
    read = DocxReader(str(files[0]))
    assert (
        read.sections[0]
        .header.paragraphs[0]
        .text.startswith("Example High · Computing   ·   Sheet")
    )
    assert "Computing Department" in read.sections[0].footer.paragraphs[0].text
    text = PdfReader(str(files[1])).pages[0].extract_text()
    assert "Example High" in text and "Computing Department" in text and "Page 1" in text


def test_house_style_from_config() -> None:
    cfg = load_config()
    assert house_style_from(cfg) == HouseStyle()
    cfg.general.house_header = "X"
    cfg.general.house_footer = "Y"
    assert house_style_from(cfg) == HouseStyle("X", "Y")
    assert house_style_from(object()) == HouseStyle()
