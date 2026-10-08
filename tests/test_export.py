"""House style: every renderer produces a readable file from the same model;
the key splits off cleanly."""

from datetime import date
from pathlib import Path

import pytest
from docx import Document as DocxReader
from pypdf import PdfReader

from staffroom.core.export import (
    AnswerGrid,
    BoxedBlocks,
    Code,
    Document,
    Heading,
    HintBox,
    Numbered,
    PageBreak,
    Paragraph,
    TableBlock,
    export,
    render,
    to_text,
)
from staffroom.core.export.pdf import box_columns


@pytest.fixture
def doc() -> Document:
    d = Document(
        title="Worksheet: Loops", topic="Python", date=date(2026, 10, 7), subtitle="seed 7"
    )
    s = d.section()
    s.add(Paragraph("Answer all questions.", style="instruction"))
    s.add(Heading("Part A"))
    s.add(Numbered(("What is a loop?", "Write a while loop."), answer_lines=2))
    s.add(Code("while x < 3:\n    x += 1", caption="Listing 1"))
    s.add(TableBlock(("A", "B", "Q"), (("0", "0", ""), ("0", "1", "")), monospace=True))
    s.add(HintBox("Count from zero."))
    s.add(BoxedBlocks(("x = 1", "while x < 3:\n    x += 1"), columns=2))
    s.add(AnswerGrid(10, per_row=8))
    s.add(PageBreak())
    k = d.section("Answer key", answers=True)
    k.add(Numbered(("Repetition.", "see listing")))
    return d


def test_sheet_and_key_split(doc: Document) -> None:
    assert doc.has_answers
    assert [s.answers for s in doc.sheet().sections] == [False]
    key = doc.key()
    assert key.title.endswith("Answer key") and [s.title for s in key.sections] == ["Answer key"]


def test_text_rendering(doc: Document) -> None:
    text = to_text(doc)
    assert text.startswith(
        "Worksheet: Loops\n================\nPython   ·   07 October 2026   ·   seed 7"
    )
    assert "  1. What is a loop?" in text and "____" in text
    assert "    while x < 3:" in text
    assert "A  B  Q" in text and "[Hint] Count from zero." in text
    assert "ANSWER KEY" in text


def test_text_boxes_and_answer_grid(doc: Document) -> None:
    text = to_text(doc)
    assert "+- 1 ---+\n| x = 1 |\n+-------+" in text
    assert "+- 2 ----------+\n| while x < 3: |\n|     x += 1   |\n+--------------+" in text
    grid = text[text.index("+--------------+") :]
    head = " ".join(str(p).center(6) for p in range(1, 9))
    assert f"{head}\n" + " ".join(["[____]"] * 8) in grid
    assert f"{'9'.center(6)} {'10'.center(6)}\n[____] [____]" in grid


def test_docx_round_trip(doc: Document, tmp_path: Path) -> None:
    path = render(doc.sheet(), tmp_path / "sheet.docx")
    read = DocxReader(str(path))
    texts = [p.text for p in read.paragraphs]
    assert "Worksheet: Loops" in texts
    assert any("1.  What is a loop?" in t for t in texts)
    assert any("while x < 3:" in t for t in texts)
    header = read.sections[0].header.paragraphs[0].text
    assert "Worksheet: Loops" in header and "07 October 2026" in header
    footer_xml = read.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml and "NUMPAGES" in footer_xml
    assert read.tables and read.tables[0].rows[0].cells[2].text == "Q"
    boxes = read.tables[2]
    assert boxes.cell(0, 0).text == "1  x = 1"
    assert boxes.cell(0, 1).text == "2  while x < 3:\n    x += 1"
    assert 'w:val="dashed"' in boxes.cell(0, 0)._tc.xml
    grid = read.tables[3]
    assert [c.text for c in grid.rows[0].cells] == [str(n) for n in range(1, 9)]
    assert [c.text for c in read.tables[4].rows[0].cells] == ["9", "10"]


def test_pdf_round_trip(doc: Document, tmp_path: Path) -> None:
    path = render(doc, tmp_path / "all.pdf")
    reader = PdfReader(str(path))
    assert len(reader.pages) >= 2  # page break + key
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Worksheet: Loops" in text and "What is a loop?" in text
    assert "while x < 3:" in text and "Count from zero." in text
    assert "Page 1" in text and "Answer key" in text
    assert "x += 1" in text  # the boxed block's second line


def test_boxes_use_fewer_columns_when_a_line_is_wide() -> None:
    short = BoxedBlocks(("x = 1", "y = 2", "z = 3"), columns=3)
    assert box_columns(short) == 3
    assert box_columns(BoxedBlocks(("x = 1",), columns=2)) == 2
    wide = BoxedBlocks(("print(" + "a, " * 20 + "a)", "y = 2"), columns=2)
    assert box_columns(wide) == 1


def test_empty_boxes_and_grid_render_nothing(tmp_path: Path) -> None:
    d = Document(title="Empty")
    d.section().add(BoxedBlocks(())).add(AnswerGrid(0))
    assert to_text(d) == "Empty\n=====\n"
    path = render(d, tmp_path / "empty.pdf")
    assert len(PdfReader(str(path)).pages) == 1
    render(d, tmp_path / "empty.docx")


def test_export_writes_sheet_and_key_per_format(doc: Document, tmp_path: Path) -> None:
    written = export(doc, tmp_path, "loops")
    assert [p.name for p in written] == [
        "loops.docx",
        "loops_key.docx",
        "loops.pdf",
        "loops_key.pdf",
    ]
    only = export(doc, tmp_path, "loops2", formats=("pdf",), split_key=False)
    assert [p.name for p in only] == ["loops2.pdf"]
    with pytest.raises(ValueError, match="unsupported"):
        render(doc, tmp_path / "x.odt")
